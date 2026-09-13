# -*- coding: utf-8 -*-
"""测试公共工具：造 .docx、跑 CLI、抓 JSON。

设计原则（为了"换设备也能跑"）：
* **不写死任何机器相关路径**——项目根 = 本文件所在目录的上一级；
* 造出来的 .docx 一律放临时目录，测试结束自动清理，不污染项目目录；
* 需要 python-docx 的用例统一先调 skip_if_no_docx()，缺依赖时"跳过"而非报错；
* 只用标准库 + 可选 python-docx，因此同一套测试在 pytest 与
  `python tests/run_tests.py`（无 pytest）下都能跑。
"""
from __future__ import annotations

import importlib.util
import io
import json
import os
import struct
import subprocess
import sys
import unittest
import zlib
from contextlib import redirect_stdout
from pathlib import Path

TESTS_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TESTS_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None


def skip_if_no_docx() -> None:
    """缺少 python-docx 时跳过（unittest.SkipTest 会被 pytest 识别为 skip）。"""
    if not DOCX_AVAILABLE:
        raise unittest.SkipTest("未安装 python-docx，跳过需要真实 .docx 的用例")


# --------------------------------------------------------------------------
# 造 .docx
# --------------------------------------------------------------------------

def _apply_run_format(run, fmt: dict) -> None:
    from docx.shared import Pt

    if "bold" in fmt:
        run.bold = fmt["bold"]
    if "italic" in fmt:
        run.italic = fmt["italic"]
    if "underline" in fmt:
        run.underline = fmt["underline"]
    if "size" in fmt:
        run.font.size = Pt(fmt["size"])


def _add_comment(doc, paragraph, text: str, author: str = "测试") -> None:
    if not hasattr(doc, "add_comment"):
        raise unittest.SkipTest("当前 python-docx 不支持批注 API（需 >= 1.2）")
    runs = paragraph.runs or [paragraph.add_run("")]
    doc.add_comment(runs=runs, text=text, author=author)


def make_docx(path, blocks, picture: str | None = None) -> str:
    """按"段落规格"造一份 .docx 并返回路径。

    blocks 元素写法：
      "纯文本"                                -> 一个普通段落
      {"runs": [("文字", {"bold": True})]}    -> 一个段落内多个 run（可分别设格式）
      {"runs": [...], "comment": "批注内容"}  -> 段落上挂一条批注

    picture：可选的图片文件路径，给出则追加到文档末尾（用于图片差异用例）。
    """
    skip_if_no_docx()
    from docx import Document

    doc = Document()
    for block in blocks:
        if isinstance(block, str):
            doc.add_paragraph(block)
            continue
        p = doc.add_paragraph()
        runs = block.get("runs")
        if runs is None:
            p.add_run(block.get("text", ""))
        for text, fmt in runs or []:
            run = p.add_run(text)
            _apply_run_format(run, fmt or {})
        if block.get("comment"):
            _add_comment(doc, p, block["comment"], block.get("author", "测试"))
    if picture:
        doc.add_picture(str(picture))
    doc.save(str(path))
    return str(path)


def make_png(path, rgb=(255, 0, 0), size=2) -> str:
    """用标准库生成一张最小合法 PNG（给"图片差异"用例当素材，免去二进制资源文件）。"""
    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + bytes(rgb) * size for _ in range(size))
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    Path(path).write_bytes(png)
    return str(path)


def intent_doc_pair(dirpath) -> tuple[str, str]:
    """造一对"有代表性"的文档，一份即覆盖主要差异类型：

    * 段落被修改 / 新增 / 删除
    * 格式变化（粗体 True→False、字号 14→12）
    * 新增一个空段落
    * 新增一条批注
    """
    base = Path(dirpath)
    a = make_docx(
        base / "v1.docx",
        [
            {"runs": [("标题", {"bold": True, "size": 14})]},
            "第一段：原文内容",
            "将被删除的段落",
        ],
    )
    b = make_docx(
        base / "v2.docx",
        [
            {"runs": [("标题", {"bold": False, "size": 12})]},
            "",
            {"runs": [("第一段：原文内容（已修改）", {})], "comment": "这里是批注"},
            "新增段落",
        ],
    )
    return a, b


# --------------------------------------------------------------------------
# 跑 CLI / 跑独立脚本
# --------------------------------------------------------------------------

def run_cli(argv) -> tuple[int, dict | None, str]:
    """在进程内调用 word_diff.cli.main(argv)，返回 (退出码, 解析出的 JSON, 原始输出)。

    同时兼容 argparse 的 SystemExit（--version、参数不足等）。
    """
    from word_diff.cli import main as cli_main

    buf = io.StringIO()
    code = 0
    with redirect_stdout(buf):
        try:
            code = cli_main([str(x) for x in argv])
        except SystemExit as exc:  # argparse 主动退出
            code = exc.code if isinstance(exc.code, int) else 1
    out = buf.getvalue()
    try:
        data = json.loads(out) if out.strip() else None
    except json.JSONDecodeError:
        data = None
    return code, data, out


def run_script(script_name: str, args) -> tuple[int, dict | None, str, str]:
    """用当前解释器把项目根下的入口脚本当"独立 CLI"运行。

    这条链路对应验收要求：**底层 Script/CLI 应能脱离智能体独立运行**。
    返回 (退出码, JSON, stdout, stderr)。
    """
    cmd = [sys.executable, str(PROJECT_ROOT / script_name), *[str(a) for a in args]]
    proc = subprocess.run(
        cmd,
        cwd=str(PROJECT_ROOT),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    try:
        data = json.loads(proc.stdout)
    except Exception:
        data = None
    return proc.returncode, data, proc.stdout, proc.stderr


def run_advanced(doc1, doc2, similarity: float | None = None) -> dict:
    """进程内调用进阶分析，返回结果字典。"""
    from word_diff.advanced import analyze_advanced

    if similarity is None:
        return analyze_advanced(str(doc1), str(doc2))
    return analyze_advanced(str(doc1), str(doc2), similarity)
